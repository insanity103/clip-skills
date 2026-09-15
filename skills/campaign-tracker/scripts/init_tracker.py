#!/usr/bin/env python3
"""Create a new posting tracker workbook for a campaign, seeded from campaign.json's posting rules.

  init_tracker.py --campaign campaign.json --out tracker.xlsx

One workbook per campaign. Add posts to it with log_post.py as they go out - this script only
creates the empty sheet and records the campaign's rules so log_post.py can check against them.
"""
import argparse
import json
import os
import sys

def die(msg):
    sys.exit(f"init_tracker: {msg}")


HEADERS = ["Clip ID", "Platform", "URL", "Posted Date", "Live-Until Date", "Views", "Likes",
          "Engagement Rate", "Min Required", "Status", "Notes"]


def build(campaign, out_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        die("openpyxl not installed - pip install openpyxl")

    wb = Workbook()
    posts = wb.active
    posts.title = "Posts"
    posts.append(HEADERS)
    for cell in posts[1]:
        cell.font = Font(bold=True)
    widths = [10, 10, 40, 14, 14, 10, 10, 14, 12, 10, 30]
    for i, w in enumerate(widths, 1):
        posts.column_dimensions[posts.cell(row=1, column=i).column_letter].width = w

    info = wb.create_sheet("Campaign")
    posting = campaign.get("posting", {})
    rows = [
        ("Name", campaign.get("name", "")),
        ("Platforms", ", ".join(campaign.get("platforms", []))),
        ("Min days live", posting.get("min_days_live", "")),
        ("Likes must be visible", posting.get("likes_visible", "")),
        ("Max reposts of the same clip", posting.get("max_reposts", "")),
        ("Paid boosting allowed", posting.get("paid_boosting", False)),
        ("Min engagement rate", posting.get("min_engagement_rate", "")),
    ]
    for r in rows:
        info.append(r)
    for cell in info["A"]:
        cell.font = Font(bold=True)
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 40

    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--campaign", required=True, help="campaign.json")
    ap.add_argument("--out", required=True, help="tracker .xlsx to create")
    ap.add_argument("--force", action="store_true", help="overwrite an existing tracker")
    args = ap.parse_args()

    if not os.path.exists(args.campaign):
        die(f"campaign file not found: {args.campaign}")
    if os.path.exists(args.out) and not args.force:
        die(f"{args.out} already exists - pass --force to overwrite, or use log_post.py to add to it")

    campaign = json.load(open(args.campaign))
    os.makedirs(os.path.dirname(os.path.abspath(args.out)) or ".", exist_ok=True)
    build(campaign, args.out)
    print(f"{args.out}: tracker created for \"{campaign.get('name', '(unnamed campaign)')}\"")


if __name__ == "__main__":
    main()
